import io
import pytest
from PIL import Image
from docx import Document
from fastapi.testclient import TestClient

from backend.app.parser.document_parser import parse_docx_bytes, parse_image_bytes
from backend.app.main import app

client = TestClient(app)

def test_parse_docx_bytes():
    doc = Document()
    doc.add_heading("Product Launch Announcement", 0)
    doc.add_paragraph("Lumina Glow Serum is now available for all members.")
    bio = io.BytesIO()
    doc.save(bio)
    file_bytes = bio.getvalue()

    result = parse_docx_bytes(file_bytes)
    assert "Lumina Glow Serum" in result["text"]
    assert result["paragraphs_count"] >= 1
    assert result["file_type"] == "docx"

def test_parse_image_bytes():
    # Create a 200x200 test RGB image with Lumina brand teal (#0f8b8d)
    img = Image.new("RGB", (200, 200), color=(15, 139, 141))
    bio = io.BytesIO()
    img.save(bio, format="PNG")
    file_bytes = bio.getvalue()

    result = parse_image_bytes(file_bytes)
    assert result["width"] == 200
    assert result["height"] == 200
    assert result["has_image"] is True
    assert len(result["dominant_colors"]) > 0

def test_upload_content_api():
    creator_headers = {"X-User-Role": "creator"}
    doc = Document()
    doc.add_paragraph("Get 100% guaranteed results with our #1 serum!")
    bio = io.BytesIO()
    doc.save(bio)

    files = {"file": ("campaign_draft.docx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {"title": "Campaign Upload Test", "content_type": "Social post"}

    resp = client.post("/content/upload", files=files, data=data, headers=creator_headers)
    assert resp.status_code == 200
    res_data = resp.json()
    assert res_data["status"] == "uploaded"
    assert "guaranteed" in res_data["extracted_text"]
